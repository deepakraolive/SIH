import os
import sys

# Ensure local lib is in python path
sys.path.insert(0, os.path.abspath('lib'))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Clean, professional palette inspired by 'Your paragraph text.pdf'
    COLOR_BG = RGBColor(255, 255, 255)
    COLOR_DARK = RGBColor(15, 23, 42)          # #0F172A Dark Slate
    COLOR_BLUE = RGBColor(0, 132, 209)         # #0084D1 Vibrant Electric Blue
    COLOR_BLUE_DARK = RGBColor(3, 105, 161)    # #0369A1 Deep Blue
    COLOR_MUTED = RGBColor(71, 85, 105)        # #475569 Slate Grey
    COLOR_CARD_BG = RGBColor(248, 250, 252)    # #F8FAFC Clean Surface
    COLOR_BORDER = RGBColor(226, 232, 240)     # #E2E8F0 Subtle Border
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_TAG_BG = RGBColor(224, 242, 254)     # #E0F2FE Sky Pill
    COLOR_GREEN = RGBColor(16, 185, 129)       # #10B981 Emerald
    COLOR_GREEN_DARK = RGBColor(5, 150, 105)   # #059669
    COLOR_RED_BG = RGBColor(254, 242, 242)
    COLOR_RED_BORDER = RGBColor(254, 202, 202)
    COLOR_RED_TEXT = RGBColor(185, 28, 28)

    blank_layout = prs.slide_layouts[6]

    def add_slide_header(slide, title_text, subtitle_text, slide_num, total_slides=7):
        # Top Meta Pill
        tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.7), Inches(0.32))
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_TAG_BG
        tag.line.color.rgb = COLOR_TAG_BG
        tf_t = tag.text_frame
        tf_t.word_wrap = False
        p = tf_t.paragraphs[0]
        p.text = "SIH 2026 · PS ID: 26091"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_DARK
        p.alignment = PP_ALIGN.CENTER

        # Right Meta
        meta_box = slide.shapes.add_textbox(Inches(7.5), Inches(0.4), Inches(5.0), Inches(0.32))
        pm = meta_box.text_frame.paragraphs[0]
        pm.text = "GramVest | Team GramVest"
        pm.font.size = Pt(10)
        pm.font.color.rgb = COLOR_MUTED
        pm.alignment = PP_ALIGN.RIGHT

        # Exact Slide Heading requested
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.65))
        pt = title_box.text_frame.paragraphs[0]
        pt.text = title_text
        pt.font.size = Pt(28)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_DARK

        # Subtitle
        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.42), Inches(11.7), Inches(0.45))
            ps = sub_box.text_frame.paragraphs[0]
            ps.text = subtitle_text
            ps.font.size = Pt(12)
            ps.font.color.rgb = COLOR_MUTED

        # Footer
        foot_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
        pf = foot_box.text_frame.paragraphs[0]
        pf.text = f"GramVest: The Rural Decision Engine                                                                                                   Slide {slide_num} of {total_slides}"
        pf.font.size = Pt(9.5)
        pf.font.color.rgb = COLOR_MUTED

    # =========================================================================
    # SLIDE 1: Title Slide (Heading: Title)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)

    # Top Pill
    t1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.55), Inches(3.4), Inches(0.36))
    t1.fill.solid()
    t1.fill.fore_color.rgb = COLOR_DARK
    t1.line.color.rgb = COLOR_DARK
    p = t1.text_frame.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Brand Title: GramVest
    brand_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(7.4), Inches(1.1))
    pb = brand_box.text_frame.paragraphs[0]
    r1 = pb.add_run()
    r1.text = "Gram"
    r1.font.size = Pt(56)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_DARK
    r2 = pb.add_run()
    r2.text = "Vest"
    r2.font.size = Pt(56)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_BLUE

    # Tagline
    tag_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(7.4), Inches(0.5))
    pt = tag_box.text_frame.paragraphs[0]
    pt.text = "The Rural Decision Engine"
    pt.font.size = Pt(22)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_BLUE

    # Hook Quote
    quote_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.75), Inches(7.4), Inches(0.6))
    pq = quote_box.text_frame.paragraphs[0]
    pq.text = '"Know your market, your competition, and your capital — before you invest a single rupee."'
    pq.font.size = Pt(13.5)
    pq.font.italic = True
    pq.font.color.rgb = COLOR_MUTED

    # Metadata Card
    meta_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.6), Inches(7.2), Inches(3.1))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = COLOR_CARD_BG
    meta_card.line.color.rgb = COLOR_BORDER
    mtf = meta_card.text_frame
    mtf.word_wrap = True

    p = mtf.paragraphs[0]
    p.text = "Problem Statement:"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    p = mtf.add_paragraph()
    p.text = "AI-Driven Hyper-Local Business Advisory and Financial Structuring Assistant for Rural Micro-Entrepreneurs"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.space_after = Pt(10)

    p = mtf.add_paragraph()
    p.text = "Theme: Rural Entrepreneurship & Market Intelligence | MSME & Agri-Business"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(4)

    p = mtf.add_paragraph()
    p.text = "Problem Statement ID: 26091 (Software Edition)"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(4)

    p = mtf.add_paragraph()
    p.text = "Presented By: Team GramVest | Round 1 Evaluation"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_MUTED

    # Right Card: Core Capabilities
    r_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.4), Inches(1.1), Inches(4.1), Inches(5.6))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = COLOR_CARD_BG
    r_card.line.color.rgb = COLOR_BLUE
    rtf = r_card.text_frame
    rtf.word_wrap = True

    p = rtf.paragraphs[0]
    p.text = "CORE INNOVATION: DUAL-ENGINE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.space_after = Pt(12)

    features = [
        ("Hyper-Local 5–10km Catchment", "Evaluates consumer base and distribution channels around the village."),
        ("Competitor Density Radar", "Identifies existing businesses to prevent sudden local market saturation."),
        ("10% / 90% Scheme Structuring", "Automatically calculates Project Cost (Margin/10%) and 90% loan eligibility."),
        ("Micro Finance vs. Term Loan Router", "Deterministic routing: Micro Finance (6.5%) vs Term Loan (8.0%)."),
        ("Quarterly EMI & Moratorium Math", "Computes debt schedule with 3-month or 6-month moratorium."),
        ("Bank-Ready DPR Dossier", "One-click generation of Detailed Project Report for bank sanction.")
    ]
    for ft, fd in features:
        p = rtf.add_paragraph()
        p.text = "✔  " + ft
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK

        p2 = rtf.add_paragraph()
        p2.text = "     " + fd
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_MUTED
        p2.space_after = Pt(5)

    # =========================================================================
    # SLIDE 2: Problem Statement (Heading: Problem Statement)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, "Problem Statement", 
                     "State Channelizing Agencies provide 90% concessional credit with only 10% margin, but rural micro-enterprises face high failure rates due to lack of local market visibility and financial illiteracy.", 2)

    cards = [
        ("01", "No Local Market Visibility", 
         "• Urban research tools ignore villages.\n• Prospective entrepreneurs cannot estimate 5–10km catchment demand.\n• Lack of data on purchasing power and fair local pricing."),
        ("02", "Invisible Competition", 
         "• Businesses chosen based on anecdotal hearsay.\n• Multiple identical copycat shops open on the same lane.\n• Triggers instant market saturation, price wars, and collapse."),
        ("03", "Financing is Confusing", 
         "• 10% equity / 90% debt calculations are intimidating.\n• Confusion between Micro Finance (≤ ₹1.4L @ 6.5%) vs. Term Loan (≤ ₹50L @ 8%).\n• Moratorium math and bank DPR filing are major hurdles."),
        ("04", "High Failure, Zero Guidance", 
         "• Over 60% of rural micro-enterprises fail in Year 1.\n• Private business consultancies charge ₹15,000–₹50,000+.\n• Zero accessible, affordable guidance for marginalized youth.")
    ]

    for i, (num, heading, desc) in enumerate(cards):
        left_pos = Inches(0.8 + i * 2.95)
        c = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.0), Inches(2.8), Inches(4.35))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD_BG
        c.line.color.rgb = COLOR_BORDER
        ctf = c.text_frame
        ctf.word_wrap = True

        p_num = ctf.paragraphs[0]
        p_num.text = num
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_BLUE
        p_num.space_after = Pt(8)

        p_h = ctf.add_paragraph()
        p_h.text = heading
        p_h.font.size = Pt(13.5)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_DARK
        p_h.space_after = Pt(10)

        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_MUTED

    # Bottom Callout
    call = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.52), Inches(11.7), Inches(0.35))
    call.fill.solid()
    call.fill.fore_color.rgb = COLOR_RED_BG
    call.line.color.rgb = COLOR_RED_BORDER
    p_c = call.text_frame.paragraphs[0]
    p_c.text = "CORE BOTTLENECK: Capital is available via SCAs and CAs, but institutional-grade feasibility intelligence has never reached the village doorstep."
    p_c.font.size = Pt(10)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_RED_TEXT
    p_c.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: Solution (Heading: Solution)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, "Solution", 
                     "GramVest turns 3 basic user inputs into a comprehensive hyper-local feasibility report and bank-ready financial roadmap in under 5 minutes.", 3)

    # 3-Step Flow Banner
    f_bar = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.98), Inches(11.7), Inches(0.72))
    f_bar.fill.solid()
    f_bar.fill.fore_color.rgb = COLOR_CARD_BG
    f_bar.line.color.rgb = COLOR_BORDER
    pf = f_bar.text_frame.paragraphs[0]
    pf.text = "STEP 1: User Inputs (Village, Margin Capital, Trade)   ➔   STEP 2: Dual-Engine Spatial & Financial Analysis   ➔   STEP 3: Bank-Ready Verdict & DPR"
    pf.font.size = Pt(11.5)
    pf.font.bold = True
    pf.font.color.rgb = COLOR_DARK
    pf.alignment = PP_ALIGN.CENTER

    # Module 1 Card
    m1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.85), Inches(5.7), Inches(4.0))
    m1.fill.solid()
    m1.fill.fore_color.rgb = COLOR_WHITE
    m1.line.color.rgb = COLOR_BLUE
    m1_tf = m1.text_frame
    m1_tf.word_wrap = True

    p = m1_tf.paragraphs[0]
    p.text = "MODULE 1: HYPER-LOCAL FEASIBILITY REPORT"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.space_after = Pt(10)

    m1_items = [
        ("Market Reach", "Estimates 5–10km radial consumer base and primary distribution channels."),
        ("Opportunity Analysis", "Highlights unserved or underserved niches in the chosen sector."),
        ("Evidence-Based SWOT", "Objective strengths, weaknesses, opportunities & threats tailored to budget."),
        ("Threats Identification", "Pinpoints supply bottlenecks, seasonal swings, and single-buyer risk."),
        ("Competitor Mapping", "Uses demographic & OSM data to estimate density of existing similar shops."),
        ("Product Market Value", "Suggests optimal pricing strategies matching regional purchasing power.")
    ]
    for t, d in m1_items:
        p = m1_tf.add_paragraph()
        p.text = f"•  {t}: {d}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # Module 2 Card
    m2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.85), Inches(5.7), Inches(4.0))
    m2.fill.solid()
    m2.fill.fore_color.rgb = COLOR_WHITE
    m2.line.color.rgb = COLOR_GREEN
    m2_tf = m2.text_frame
    m2_tf.word_wrap = True

    p = m2_tf.paragraphs[0]
    p.text = "MODULE 2: SMART SCHEME CALCULATOR & ROUTER"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN_DARK
    p.space_after = Pt(10)

    m2_items = [
        ("Financial Structuring", "Project Cost = Margin / 10%. Max Loan = 90% of Project Cost. (e.g. ₹1,00,000 margin enables ₹10,00,000 project & ₹9,00,000 loan)."),
        ("Scheme Auto-Selection", "Deterministic rule-based routing to the exact concessional loan:"),
        ("   - Micro Finance Scheme", "Project Cost ≤ ₹1.40 Lakh: 6.5% interest p.a., 3-yr tenure, 3-mo moratorium."),
        ("   - Term Loan Scheme", "Project Cost ₹1.40L to ₹50 Lakh: 8.0% interest p.a., 7-yr tenure, 6-mo moratorium."),
        ("EMI & Moratorium Schedule", "Generates quarterly repayment schedules, operational costs, and working capital."),
        ("Bank-Ready DPR Dossier", "One-click Detailed Project Report download for immediate bank appraisal.")
    ]
    for t, d in m2_items:
        p = m2_tf.add_paragraph()
        p.text = f"•  {t}: {d}" if not t.startswith("   -") else f"    {t}: {d}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 4: Flow Chart (Heading: Flow Chart)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, "Flow Chart", 
                     "End-to-end operational decision flow from user inputs to bank-sanctionable feasibility report.", 4)

    flowchart_path = "FlowChart.jpeg"
    if os.path.exists(flowchart_path):
        s4.shapes.add_picture(flowchart_path, Inches(0.8), Inches(1.98), width=Inches(5.6))
    else:
        box = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.98), Inches(5.6), Inches(4.8))
        box.text_frame.text = "[FlowChart.jpeg]"

    # Right Box: Step Details
    right_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.98), Inches(5.8), Inches(4.85))
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = COLOR_CARD_BG
    right_box.line.color.rgb = COLOR_BORDER
    r_tf = right_box.text_frame
    r_tf.word_wrap = True

    p = r_tf.paragraphs[0]
    p.text = "BUSINESS ANALYSIS & DECISION FLOW"
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK
    p.space_after = Pt(10)

    steps = [
        ("1. User Input Layer", "Collects village location, available margin capital, and proposed business."),
        ("2. Data Engine & GIS Ingestion", "Pulls 5–10km catchment, demographics, and mandi commodity rates."),
        ("3. Feasibility & Threat Engine", "Evaluates market saturation, competitor density, and evidence-based SWOT."),
        ("4. Deterministic Financial Engine", "Calculates Project Cost, 90% Loan, quarterly EMI, cash flow, and DSCR."),
        ("5. Scheme Router", "Matches Micro Finance (≤ ₹1.4L @ 6.5%) or Term Loan (≤ ₹50L @ 8%)."),
        ("6. Grounded AI Advisor & RAG", "Explains insights in local vernacular via RAG over official scheme gazettes."),
        ("7. Final DPR & Action Roadmap", "Generates official bank-ready Detailed Project Report for loan sanctioning.")
    ]
    for st, d in steps:
        p = r_tf.add_paragraph()
        p.text = f"✔  {st}"
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE

        p2 = r_tf.add_paragraph()
        p2.text = f"     {d}"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_MUTED
        p2.space_after = Pt(4)

    # =========================================================================
    # SLIDE 5: Tech Stack (Heading: Tech Stack)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, "Tech Stack", 
                     "Deterministic software core + grounded AI layer. Golden Engineering Rule: 'AI explains; software calculates.' Zero financial hallucinations.", 5)

    tech_cols = [
        ("Frontend & UI", "React.js & Tailwind CSS", [
            ("Mobile-First Design", "Responsive dashboard optimized for low-bandwidth rural smartphone users."),
            ("GIS Catchment Maps", "Leaflet & Mapbox for 5–10 km radial catchment boundary visualization."),
            ("Guided Setup Wizard", "Simple 3-step input requiring zero prior business experience."),
            ("Multilingual Ready", "Interface architecture built for instant vernacular translation.")
        ]),
        ("Backend & Financials", "FastAPI & Python Math Core", [
            ("Deterministic Finance", "Python Decimal & NumPy for 100% audit-proof loan, EMI, and DSCR math."),
            ("Versioned Scheme Rules", "Policy-driven JSON rules for Micro Finance and Term Loan parameters."),
            ("Asynchronous APIs", "FastAPI backend handles concurrent requests with sub-second latency."),
            ("DPR PDF Engine", "ReportLab integration to generate official, bank-ready PDF dossiers.")
        ]),
        ("AI, GIS & Data Layer", "PostgreSQL, PostGIS & RAG", [
            ("Geospatial Database", "PostgreSQL + PostGIS for spatial queries and competitor clustering."),
            ("Semantic Vector Search", "pgvector and FAISS for RAG over verified official scheme gazettes."),
            ("Numerical Guardrails", "Strict verification layer ensures LLM never invents loan numbers or rates."),
            ("External Feeds", "OpenStreetMap Overpass API and Agmarknet Mandi wholesale price feeds.")
        ])
    ]

    for idx, (title, stack, items) in enumerate(tech_cols):
        cx = Inches(0.8 + idx * 3.95)
        c = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.98), Inches(3.8), Inches(4.4))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD_BG
        c.line.color.rgb = COLOR_BORDER
        ctf = c.text_frame
        ctf.word_wrap = True

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK

        p_s = ctf.add_paragraph()
        p_s.text = stack
        p_s.font.size = Pt(10)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_BLUE
        p_s.space_after = Pt(12)

        for it, idesc in items:
            p_it = ctf.add_paragraph()
            p_it.text = f"•  {it}:"
            p_it.font.size = Pt(10)
            p_it.font.bold = True
            p_it.font.color.rgb = COLOR_DARK

            p_id = ctf.add_paragraph()
            p_id.text = f"   {idesc}"
            p_id.font.size = Pt(9)
            p_id.font.color.rgb = COLOR_MUTED
            p_id.space_after = Pt(4)

    # Bottom Banner
    bb = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.52), Inches(11.7), Inches(0.35))
    bb.fill.solid()
    bb.fill.fore_color.rgb = COLOR_DARK
    bb.line.color.rgb = COLOR_DARK
    p_b = bb.text_frame.paragraphs[0]
    p_b.text = "CORE ENGINEERING PRINCIPLE: 'AI explains; software calculates.' Generative AI never computes financial values."
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = RGBColor(56, 189, 248)
    p_b.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 6: Future Scope (Heading: Future Scope)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, "Future Scope", 
                     "Strategic roadmap to scale GramVest from a hackathon prototype into national digital public infrastructure.", 6)

    scope_cards = [
        ("Phase 1: Near Term", "Vernacular Voice AI (Bhashini)", 
         "• Integration with GoI's Bhashini voice engine.\n• Voice-in and voice-out in Hindi, Punjabi, Marathi, Tamil & Bengali.\n• Removes digital & textual literacy barriers for rural artisans and women SHGs."),
        
        ("Phase 2: Strategic", "Direct Bank & SCA API Integration", 
         "• Direct digital connectivity with State Channelizing Agencies (SCAs) & JanSamarth.\n• Enables 1-click digital loan application submission.\n• Real-time digital sanction and disbursement tracking."),
        
        ("Phase 3: Intelligence", "Satellite & IoT Agro Data", 
         "• Ingesting ISRO Bhuvan & Copernicus satellite imagery.\n• Real-time soil health, crop cycles, and mandi demand spikes.\n• Powers precision agro-processing feasibility models."),
        
        ("Phase 4: Sustainability", "Post-Sanction WhatsApp Bot", 
         "• Continuous lightweight WhatsApp micro-advisory.\n• Quarterly repayment reminders and cash-flow health checks.\n• Seasonal price warnings to actively prevent business defaults.")
    ]

    for i, (ph, title, desc) in enumerate(scope_cards):
        left_pos = Inches(0.8 + i * 2.95)
        c = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.0), Inches(2.8), Inches(4.35))
        c.fill.solid()
        c.fill.fore_color.rgb = COLOR_CARD_BG
        c.line.color.rgb = COLOR_BORDER
        ctf = c.text_frame
        ctf.word_wrap = True

        p_ph = ctf.paragraphs[0]
        p_ph.text = ph
        p_ph.font.size = Pt(10)
        p_ph.font.bold = True
        p_ph.font.color.rgb = COLOR_BLUE
        p_ph.space_after = Pt(6)

        p_t = ctf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_DARK
        p_t.space_after = Pt(10)

        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = COLOR_MUTED

    # Bottom Banner
    sum_b = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.52), Inches(11.7), Inches(0.35))
    sum_b.fill.solid()
    sum_b.fill.fore_color.rgb = RGBColor(236, 253, 245)
    sum_b.line.color.rgb = RGBColor(167, 243, 208)
    p_s = sum_b.text_frame.paragraphs[0]
    p_s.text = "MISSION: Transforming rural entrepreneurship from anecdotal guesswork into data-driven conviction across 600,000+ Indian villages."
    p_s.font.size = Pt(10)
    p_s.font.bold = True
    p_s.font.color.rgb = RGBColor(6, 95, 70)
    p_s.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 7: Impact & Conclusion (Heading: Impact & Conclusion)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_header(s7, "Impact & Conclusion", 
                     "GramVest replaces weeks of anecdotal guesswork with minutes of institutional-grade financial clarity, empowering marginalized entrepreneurs and protecting public capital.", 7)

    impacts = [
        ("65%", "Rural Population Reach", "Designed to scale across all 6,00,000+ Indian villages at near-zero marginal cost."),
        ("Weeks ➔ 5m", "Time to Feasibility", "Slashes business feasibility assessment from 3 weeks of manual inquiry to under 5 minutes."),
        ("₹0", "Consulting Cost", "Replaces expensive ₹15,000–₹50,000 consultants with free, institutional-grade AI advisory."),
        ("90% : 10%", "Scheme Clarity", "Exact debt-equity structuring, eliminating surprise shortfalls and bank rejections.")
    ]

    for idx, (stat, title, desc) in enumerate(impacts):
        ix = Inches(0.8 + idx * 2.95)
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ix, Inches(2.0), Inches(2.8), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_BORDER
        ctf = card.text_frame
        ctf.word_wrap = True

        p_s = ctf.paragraphs[0]
        p_s.text = stat
        p_s.font.size = Pt(28)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_BLUE
        p_s.alignment = PP_ALIGN.CENTER

        p_t = ctf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_DARK
        p_t.alignment = PP_ALIGN.CENTER
        p_t.space_after = Pt(4)

        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_MUTED
        p_d.alignment = PP_ALIGN.CENTER

    # Closing banner
    banner = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.55), Inches(11.7), Inches(2.2))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_BLUE
    banner.line.color.rgb = COLOR_BLUE
    btf = banner.text_frame
    btf.word_wrap = True

    bp = btf.paragraphs[0]
    bp.text = '"Every village has a market. GramVest helps rural India see it clearly."'
    bp.font.size = Pt(20)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_WHITE
    bp.alignment = PP_ALIGN.CENTER
    bp.space_after = Pt(10)

    bp2 = btf.add_paragraph()
    bp2.text = "Ready to empower the next 10 million rural micro-enterprises with data-backed conviction."
    bp2.font.size = Pt(14)
    bp2.font.color.rgb = COLOR_WHITE
    bp2.alignment = PP_ALIGN.CENTER
    bp2.space_after = Pt(16)

    bp3 = btf.add_paragraph()
    bp3.text = "Team GramVest · Smart India Hackathon 2026 · Thank You!"
    bp3.font.size = Pt(13)
    bp3.font.bold = True
    bp3.font.color.rgb = RGBColor(224, 242, 254)
    bp3.alignment = PP_ALIGN.CENTER

    output_path = "GramVest_SIH_Round1.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path} (Size: {os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    create_presentation()
