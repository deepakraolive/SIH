"""
GramVest - SIH 2026 Round 1 Presentation Deck Generator
Generates a PowerPoint (.pptx) deck matching 'Your paragraph text.pdf' and SIH requirements.

Prerequisite:
pip install python-pptx
"""

import sys
import os

def create_deck():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        print("python-pptx is not installed. Run: pip install python-pptx")
        print("You can view and present the interactive presentation directly via 'index.html'!")
        return

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    COLOR_DARK = RGBColor(15, 23, 42)      # #0f172a
    COLOR_BLUE = RGBColor(0, 132, 209)     # #0084d1
    COLOR_MUTED = RGBColor(71, 85, 105)    # #475569
    COLOR_LIGHT = RGBColor(248, 250, 252)  # #f8fafc
    COLOR_BORDER = RGBColor(226, 232, 240) # #e2e8f0
    COLOR_GREEN = RGBColor(16, 185, 129)   # #10b981
    COLOR_WHITE = RGBColor(255, 255, 255)

    blank_layout = prs.slide_layouts[6]

    def add_top_bar(slide, tag_text, slide_num):
        # Badge
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.2), Inches(0.35))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_DARK
        shape.line.color.rgb = COLOR_DARK
        tf = shape.text_frame
        tf.text = tag_text
        tf.paragraphs[0].font.size = Pt(9.5)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = COLOR_WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Right header
        txBox = slide.shapes.add_textbox(Inches(8.0), Inches(0.4), Inches(4.5), Inches(0.35))
        p = txBox.text_frame.paragraphs[0]
        p.text = "GramVest | Smart India Hackathon 2026"
        p.font.size = Pt(10)
        p.font.color.rgb = COLOR_MUTED
        p.alignment = PP_ALIGN.RIGHT

        # Footer
        footerBox = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.35))
        p_foot = footerBox.text_frame.paragraphs[0]
        p_foot.text = f"GramVest | The Rural Decision Engine                                                                                                   Slide {slide_num} of 7"
        p_foot.font.size = Pt(9.5)
        p_foot.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # SLIDE 1: Title
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    
    # Left text box
    tb = slide1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(7.5), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "GramVest"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE

    p2 = tf.add_paragraph()
    p2.text = "The Rural Decision Engine"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_DARK
    p2.space_after = Pt(14)

    p3 = tf.add_paragraph()
    p3.text = '"Know your market, your competition, and your capital — before you invest a single rupee."'
    p3.font.size = Pt(15)
    p3.font.italic = True
    p3.font.color.rgb = COLOR_MUTED
    p3.space_after = Pt(24)

    p4 = tf.add_paragraph()
    p4.text = "Smart India Hackathon 2026 · Software Edition\nProblem Statement ID: 26091\nTheme: Rural Entrepreneurship & Market Intelligence | MSME & Agri-Business"
    p4.font.size = Pt(12)
    p4.font.color.rgb = COLOR_DARK

    # Right Card
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.5), Inches(4.0), Inches(4.8))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_LIGHT
    card1.line.color.rgb = COLOR_BLUE
    c_tf = card1.text_frame
    c_tf.word_wrap = True
    cp = c_tf.paragraphs[0]
    cp.text = "CORE INNOVATION: DUAL ENGINE"
    cp.font.size = Pt(11)
    cp.font.bold = True
    cp.font.color.rgb = COLOR_BLUE
    cp.space_after = Pt(10)

    bullets = [
        "1. Hyper-Local GIS Radar (5–10km catchment reach)",
        "2. Competitor Density & Saturation Index",
        "3. Deterministic Scheme Router (10% Equity / 90% Loan)",
        "4. Micro Finance (≤ ₹1.4L @ 6.5%) & Term Loan (≤ ₹50L @ 8%)",
        "5. Quarterly EMI & Moratorium Simulator",
        "6. One-Click Bank-Ready DPR Dossier"
    ]
    for b in bullets:
        bp = c_tf.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(11)
        bp.font.color.rgb = COLOR_DARK
        bp.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_top_bar(slide2, "THE PROBLEM · SIH 26091", 2)

    tb = slide2.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.2))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Rural entrepreneurs invest blind"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    p2 = tf.add_paragraph()
    p2.text = "SCAs provide 90% concessional credit with 10% promoter margin, but over 60% of rural ventures fail within Year 1 due to zero local market visibility and confusing financial structures."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

    cards_data = [
        ("1", "No Local Market Visibility", "Urban tools ignore villages. Entrepreneurs cannot calculate 5-10km catchment demand, seasonal pricing, or buyer purchasing power."),
        ("2", "Competition is Invisible", "Ventures rely on anecdotal hearsay. Opening multiple identical shops on one street causes sudden saturation and failure."),
        ("3", "Financing is Confusing", "Calculating 10% equity / 90% debt, Micro Finance vs. Term Loans, moratorium periods, and bank DPRs is overwhelming."),
        ("4", "High Failure, Zero Guidance", "Consultants charge ₹15,000–₹50,000+ and avoid rural blocks. Beneficiaries risk their life savings with zero guidance.")
    ]

    for idx, (num, heading, desc) in enumerate(cards_data):
        cx = Inches(0.8 + idx * 2.95)
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(2.4), Inches(2.8), Inches(4.0))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_BORDER
        ctf = card.text_frame
        ctf.word_wrap = True

        p_num = ctf.paragraphs[0]
        p_num.text = num
        p_num.font.size = Pt(16)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_BLUE

        p_h = ctf.add_paragraph()
        p_h.text = heading
        p_h.font.size = Pt(14)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_DARK
        p_h.space_after = Pt(8)

        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_top_bar(slide3, "OUR SOLUTION · GRAMVEST", 3)

    tb = slide3.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "One platform. Every decision a rural founder needs."
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    p2 = tf.add_paragraph()
    p2.text = "GramVest turns 3 simple inputs (Location, Available Margin, Business Idea) into a hyper-local feasibility report and bank-ready scheme roadmap in under 5 minutes."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

    # Module 1 Box
    m1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(5.7), Inches(4.2))
    m1.fill.solid()
    m1.fill.fore_color.rgb = COLOR_WHITE
    m1.line.color.rgb = COLOR_BLUE
    m1_tf = m1.text_frame
    m1_tf.word_wrap = True
    m1_p = m1_tf.paragraphs[0]
    m1_p.text = "MODULE 1: HYPER-LOCAL FEASIBILITY ENGINE"
    m1_p.font.size = Pt(12)
    m1_p.font.bold = True
    m1_p.font.color.rgb = COLOR_BLUE
    m1_p.space_after = Pt(10)

    m1_items = [
        "• Market Reach: 5–10km radial catchment & distribution channels",
        "• Opportunity Analysis: Unserved demand gaps and product deficits",
        "• Evidence-Based SWOT: Grounded in local economic reality",
        "• Threat Identification: Supply chain bottlenecks & seasonal swings",
        "• Competitor Mapping: Estimated density of similar local businesses",
        "• Product Market Value: Optimal pricing for rural purchasing power"
    ]
    for it in m1_items:
        p = m1_tf.add_paragraph()
        p.text = it
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # Module 2 Box
    m2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.3), Inches(5.7), Inches(4.2))
    m2.fill.solid()
    m2.fill.fore_color.rgb = COLOR_WHITE
    m2.line.color.rgb = COLOR_GREEN
    m2_tf = m2.text_frame
    m2_tf.word_wrap = True
    m2_p = m2_tf.paragraphs[0]
    m2_p.text = "MODULE 2: SMART SCHEME CALCULATOR & ROUTER"
    m2_p.font.size = Pt(12)
    m2_p.font.bold = True
    m2_p.font.color.rgb = COLOR_GREEN
    m2_p.space_after = Pt(10)

    m2_items = [
        "• Deterministic 10x Structuring: Project Cost = Margin / 10% (90% Loan)",
        "• Scheme Auto-Routing:",
        "    - Micro Finance (Cost ≤ ₹1.40L): 6.5% interest, 3-yr tenure, 3-mo moratorium",
        "    - Term Loan (Cost ₹1.40L–₹50L): 8.0% interest, 7-yr tenure, 6-mo moratorium",
        "• EMI & Moratorium Generator: Quarterly repayment schedule & cash flow",
        "• Bank-Ready DPR: One-click standard Detailed Project Report"
    ]
    for it in m2_items:
        p = m2_tf.add_paragraph()
        p.text = it
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 4: Process Flow (Flowchart)
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_top_bar(slide4, "PROCESS WORKFLOW · USER JOURNEY", 4)

    tb = slide4.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Seamless Decision Flow: From Input to Sanction"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    if os.path.exists("FlowChart.jpeg"):
        slide4.shapes.add_picture("FlowChart.jpeg", Inches(0.8), Inches(2.2), width=Inches(5.6))
    else:
        box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3))
        box.text_frame.text = "[FlowChart.jpeg]"

    # Right side steps
    steps_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.4))
    steps_box.fill.solid()
    steps_box.fill.fore_color.rgb = COLOR_LIGHT
    steps_box.line.color.rgb = COLOR_BORDER
    s_tf = steps_box.text_frame
    s_tf.word_wrap = True

    flow_steps = [
        "1. User Input: Village location, margin capital, business category",
        "2. Data Engine: Ingests 5–10km GIS, mandi prices, census density",
        "3. Feasibility Engine: Assesses competition headroom & local threats",
        "4. Financial Engine: Computes project cost, 90% debt, quarterly EMI",
        "5. Scheme Router: Micro Finance vs. Term Loan policy matching",
        "6. Grounded AI Advisor: RAG over verified government gazettes",
        "7. DPR Dossier: One-click bank-ready report for loan appraisal"
    ]
    sp = s_tf.paragraphs[0]
    sp.text = "DECISION PIPELINE BREAKDOWN"
    sp.font.size = Pt(12)
    sp.font.bold = True
    sp.font.color.rgb = COLOR_DARK
    sp.space_after = Pt(8)

    for step in flow_steps:
        p = s_tf.add_paragraph()
        p.text = "✔ " + step
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_DARK
        p.space_after = Pt(5)

    # -------------------------------------------------------------
    # SLIDE 5: Tech Stack & System Architecture
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_top_bar(slide5, "SYSTEM ARCHITECTURE · TECH STACK", 5)

    tb = slide5.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Deterministic Software Core + Grounded AI Layer"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    p2 = tf.add_paragraph()
    p2.text = "Golden Engineering Principle: AI Explains; Software Calculates. Strict numerical guardrails eliminate hallucinated loan and EMI figures."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

    t_data = [
        ("Frontend / UI", "React.js, Tailwind CSS, Leaflet / Mapbox", [
            "• Mobile-first guided setup wizard",
            "• 5–10km catchment GIS map view",
            "• Interactive What-If stress sliders",
            "• Multilingual interface (En, Hi, Pa)"
        ]),
        ("Backend & API Layer", "FastAPI, NumPy, Decimal, ReportLab", [
            "• Deterministic math for EMI, DSCR & cash flow",
            "• Versioned scheme policy rule matrix",
            "• Strict Pydantic API schemas",
            "• Bank-ready DPR PDF compilation"
        ]),
        ("AI, GIS & Storage", "PostgreSQL, PostGIS, pgvector, LangChain", [
            "• Spatial radius queries for rival density",
            "• RAG over official scheme gazettes",
            "• Numerical guardrail rejects math errors",
            "• OpenStreetMap & Mandi API feeds"
        ])
    ]

    for idx, (title, stack, bullets) in enumerate(t_data):
        tx = Inches(0.8 + idx * 3.95)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(2.3), Inches(3.8), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_BORDER
        ctf = card.text_frame
        ctf.word_wrap = True

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK

        p_st = ctf.add_paragraph()
        p_st.text = stack
        p_st.font.size = Pt(10.5)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_BLUE
        p_st.space_after = Pt(10)

        for b in bullets:
            bp = ctf.add_paragraph()
            bp.text = b
            bp.font.size = Pt(11)
            bp.font.color.rgb = COLOR_MUTED
            bp.space_after = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 6: Future Scope & Roadmap
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_top_bar(slide6, "FUTURE SCOPE · ROADMAP", 6)

    tb = slide6.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Roadmap Beyond the Hackathon"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    p2 = tf.add_paragraph()
    p2.text = "A clear 4-stage evolutionary path to transform GramVest from a hackathon prototype into national digital public infrastructure."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

    phases = [
        ("Phase 1: Near Term", "Vernacular Voice AI (Bhashini)", "Voice-in/voice-out via Bhashini API in Hindi, Punjabi, Marathi, Tamil & Bengali. Removes literacy barrier for rural artisans & SHGs."),
        ("Phase 2: Strategic", "Direct SCA & Bank API", "Direct digital tie-ins with State Channelizing Agencies and JanSamarth portal for 1-click loan submission & tracking."),
        ("Phase 3: Intelligence", "Satellite & IoT Agro Data", "Ingesting ISRO Bhuvan satellite imagery for soil health, harvest cycles, and regional mandi supply fluctuations."),
        ("Phase 4: Sustainability", "Post-Sanction WhatsApp Bot", "Ongoing micro-advisory tracking quarterly repayment nudges, raw material alerts, and default prevention.")
    ]

    for idx, (phase, title, desc) in enumerate(phases):
        px = Inches(0.8 + idx * 2.95)
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, Inches(2.3), Inches(2.8), Inches(4.3))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_BORDER
        ctf = card.text_frame
        ctf.word_wrap = True

        p_ph = ctf.paragraphs[0]
        p_ph.text = phase
        p_ph.font.size = Pt(11)
        p_ph.font.bold = True
        p_ph.font.color.rgb = COLOR_BLUE

        p_t = ctf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_DARK
        p_t.space_after = Pt(8)

        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # SLIDE 7: Impact & Conclusion
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_top_bar(slide7, "IMPACT & CONCLUSION", 7)

    tb = slide7.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "De-risking Grassroots Lending Across Bharat"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_DARK

    p2 = tf.add_paragraph()
    p2.text = "GramVest replaces weeks of anecdotal guesswork with minutes of institutional-grade financial clarity, empowering marginalized entrepreneurs and protecting public capital."
    p2.font.size = Pt(13)
    p2.font.color.rgb = COLOR_MUTED

    impacts = [
        ("65%", "Rural Population Reach", "Designed to scale across all 6,00,000+ Indian villages at near-zero marginal cost."),
        ("Weeks → 5m", "Time to Feasibility", "Slashes business evaluation time from 3 weeks of manual inquiry to under 5 minutes."),
        ("₹0", "Consulting Cost", "Replaces expensive ₹15,000–₹50,000 consultants with free, institutional AI advisory."),
        ("90% : 10%", "Scheme Clarity", "Exact debt-equity structuring, eliminating surprise shortfalls and bank rejections.")
    ]

    for idx, (stat, title, desc) in enumerate(impacts):
        ix = Inches(0.8 + idx * 2.95)
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ix, Inches(2.3), Inches(2.8), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_BORDER
        ctf = card.text_frame
        ctf.word_wrap = True

        p_s = ctf.paragraphs[0]
        p_s.text = stat
        p_s.font.size = Pt(28)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_BLUE
        p_s.alignment = PP_ALIGN.CENTER

        p_t = ctf.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_DARK
        p_t.alignment = PP_ALIGN.CENTER
        p_t.space_after = Pt(4)

        p_d = ctf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = COLOR_MUTED
        p_d.alignment = PP_ALIGN.CENTER

    # Closing banner
    banner = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_BLUE
    banner.line.color.rgb = COLOR_BLUE
    b_tf = banner.text_frame
    b_tf.word_wrap = True
    bp = b_tf.paragraphs[0]
    bp.text = '"Every village has a market. GramVest helps rural India see it clearly."'
    bp.font.size = Pt(18)
    bp.font.bold = True
    bp.font.color.rgb = COLOR_WHITE
    bp.space_after = Pt(6)

    bp2 = b_tf.add_paragraph()
    bp2.text = "Team GramVest · Smart India Hackathon 2026 · Thank You!"
    bp2.font.size = Pt(13)
    bp2.font.color.rgb = COLOR_WHITE

    output_path = "GramVest_SIH_Round1.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
